import docx
from openpyxl import Workbook

try:
    doc = docx.Document("experment_Page_0737.docx") # file to import
except:
    print('There is some error while opening the docx file')

mydata = ''

for para in doc.paragraphs:
    # print(para.text)
   # list(para.text)
   #mydata.append(para.text)
   mydata = mydata + ' ' + para.text
   # print('---------------------------')

print(mydata)
listofEntries = []

for i, elem in enumerate(str(mydata).split("~~")):
    listofEntries.append([])
    for j in str(elem).split('~'):
        data = str(j).strip()
        
        if data.find('$') != -1:
            listofEntries[i].append(" ")
            continue
        
        if data == '0':
            listofEntries[i].append(data.replace("0", " "))
        else:
            listofEntries[i].append(data.replace("\t", " "))
        
            


finalfinalList = []

for i in listofEntries:
    if len(i) < 16:
        continue
    else:
        finalfinalList.append(i)

#purifiedData = []

#for i in mydata:
#    purifiedData.append(str(i).replace("\t", " "))

#print(purifiedData)

# Converting to excel sheet--------------------------------------------------------------------------
try:
    wb = Workbook()
    # f = open("extraction" + str(i) +".txt", "r")
    # arrangedData = dataManipulation.dataManipulation('extraction' + str(i) + '.txt') #set file name here
    # print(arrangedData)
    # Get Arranged data in list of list


    ws = wb.active

    # OUR MAIN ISSUE
    for j in finalfinalList: # purified data should be like [['','','','',''],['','','','',''],['','','','','']]

      # upper case naming and replace ' ' with double spaces 
      j[3] = str(j[3]).upper()
      j[3] = j[3].replace(" ", "  ")

      # phone number to a format
      if len(j[7]) == 10:
          j[7] = '+91 ' + j[7][0:5] + ' ' + j[7][5:]
    
      if len(j[9]) == 10:
          j[9] = '+91 ' + j[9][0:5] + ' ' + j[9][5:]


      # capitalize
      j[10] = str(j[10]).capitalize()
      j[11] = str(j[11]).capitalize()
      j[15] = str(j[15]).capitalize()

      # open close upper
      j[5] = j[5].upper()
    
      
      ws.append(j)

    
    wb.save('experment_Page_0737.xlsx')
    
 
except Exception as e:
  print("An exception occurred: ", e)
  #sfks;ldfkf;l
  #sfks;ldfkf;l

  #sfks;ldfkf;l


  #sfks;ldfkf;l
